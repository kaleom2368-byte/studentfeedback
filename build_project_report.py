from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE

OUTPUT = r"D:\clg2\StudentFeedback\Student_Feedback_System_Project_Report.docx"

NAVY = "0B2545"
BLUE = "2E74B5"
MID_BLUE = "1F4D78"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GRAY = "5B6573"
RED = "9B1C1C"
GREEN = "1F6B45"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for index, width in enumerate(widths):
        grid.gridCol_lst[index].set(qn("w:w"), str(width))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_font(run, size=11, color=NAVY, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic

def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        r = p.add_run(item)
        set_font(r)

def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        r = p.add_run(item)
        set_font(r)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size={1:16, 2:13, 3:12}[level], color={1:BLUE, 2:BLUE, 3:MID_BLUE}[level], bold=True)
    return p

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    header = table.rows[0].cells
    for i, title in enumerate(headers):
        set_cell_shading(header[i], LIGHT)
        p = header[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        set_font(r, size=10, color=NAVY, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[i], PALE)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, size=9.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

def add_callout(doc, label, text, color=LIGHT):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + " ")
    set_font(r, size=10.5, color=NAVY, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(.492)
    section.footer_distance = Inches(.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for level, size, color, before, after in [(1,16,BLUE,16,8),(2,13,BLUE,12,6),(3,12,MID_BLUE,8,4)]:
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        s = styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(11)
        s.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("ANONYMOUS STUDENT FEEDBACK SYSTEM  |  PROJECT REPORT")
    set_font(r, size=8.5, color=GRAY, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Student Feedback System  |  August 2026")
    set_font(r, size=8.5, color=GRAY)
    return doc

doc = setup_document()

# Cover
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run("PROJECT REPORT")
set_font(r, size=11, color=BLUE, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Anonymous Student Feedback System")
set_font(r, size=27, color=NAVY, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(26)
r = p.add_run("A privacy-conscious web application for student feedback, analytics, and academic oversight")
set_font(r, size=13, color=GRAY)
add_callout(doc, "Report scope:", "Current architecture, implemented feedback-cycle foundation, security posture, workflows, and recommended next steps.")
for _ in range(9):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Prepared from the current project source and verified database design")
set_font(r, size=10.5, color=GRAY, italic=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("August 31, 2026")
set_font(r, size=10.5, color=GRAY)

doc.add_page_break()
add_heading(doc, "Executive Summary", 1)
add_para(doc, "The Anonymous Student Feedback System is a college-focused web application that enables students to provide feedback on faculty while giving Faculty and HOD users aggregated dashboards. The platform is designed to protect the link between a student identity and an individual feedback response.")
add_para(doc, "The current implementation uses a browser-based frontend, a Node.js/Express backend, and a MySQL database. The most recent foundation work makes feedback-cycle tracking explicit: every anonymous feedback response and each separate participation record are associated with the same cycle identifier. This eliminates date-based ambiguity and prepares the system for reliable participation reporting.")
add_callout(doc, "Current status:", "Student, Faculty, HOD, and static Admin interfaces exist. Faculty analytics and anonymous feedback workflows are functional. Feedback-cycle data integrity has been strengthened. Admin authentication and cycle-management features remain pending completion.")

add_heading(doc, "System Objectives", 2)
add_bullets(doc, [
    "Collect structured student feedback about faculty and subjects.",
    "Keep feedback content anonymous by storing it separately from student participation tracking.",
    "Provide faculty with professional analytics, ratings, trends, and anonymous comments.",
    "Provide HOD users with department-level faculty performance information.",
    "Establish an auditable feedback-cycle foundation so reporting and participation status are cycle-specific.",
])

add_heading(doc, "Technology Stack", 2)
add_table(doc, ["Layer", "Technology", "Role"], [
    ("Frontend", "HTML5, CSS3, Vanilla JavaScript", "Portal pages, forms, dashboards, charts, and client-side rendering."),
    ("Backend", "Node.js 22, Express 5", "Routes, session handling, validation, data access, and API responses."),
    ("Database", "MySQL (Aiven-hosted)", "User, faculty, anonymous feedback, cycle, and participation data."),
    ("Security", "Helmet, express-session, MySQL session store", "HTTP protection headers and server-side session persistence."),
    ("Visualization", "Chart.js", "Faculty dashboard charts and performance-trend display."),
], [1600, 2100, 5660])

doc.add_page_break()
add_heading(doc, "Architecture Overview", 1)
add_para(doc, "The application is organized as a conventional Express project. Static frontend assets are served from the public directory, while route modules provide login, dashboard, feedback, directory, and role-specific API behavior.")
add_table(doc, ["Area", "Primary Components", "Responsibility"], [
    ("Application entry", "server.js", "Creates Express app, configures security/session middleware, serves public files, and mounts route modules."),
    ("Database access", "db.js", "Creates the MySQL connection pool using environment variables."),
    ("Role routes", "routes/student.js, faculty.js, hod.js, admin.js", "Authentication and role-specific API behavior."),
    ("Feedback workflow", "routes/feedback.js", "Faculty list, anonymous feedback submission, and separate submission tracking."),
    ("Frontend", "public/auth, public/dashboard, public/js, public/css", "Login forms, portals, dashboard views, JavaScript rendering, and styles."),
    ("Migrations", "migrate-cycle-aware-submissions.js and related scripts", "Safe database evolution and historical-data handling."),
], [1900, 3200, 4260])

add_heading(doc, "Role-Based Portal Summary", 2)
add_table(doc, ["Role", "Current Capabilities", "Access Model"], [
    ("Student", "Login, select faculty, submit feedback, view own submission history/status.", "Student session stored in req.session.student."),
    ("Faculty", "Login, profile view, current-cycle feedback analytics, participation totals, anonymous comments.", "Faculty session stored in req.session.faculty."),
    ("HOD", "Login, department dashboard, faculty performance summaries, individual faculty analysis.", "HOD session stored in req.session.hod."),
    ("Admin", "Static login/dashboard pages currently exist; secure operational authorization is pending setup.", "No completed req.session.admin flow yet."),
], [1200, 5000, 3160])

add_heading(doc, "Core Data Model", 1)
add_table(doc, ["Table", "Key Fields", "Purpose"], [
    ("students", "student_id, name, department, year, division", "Stores student identities and academic grouping."),
    ("faculty", "faculty_id, name, department, subject", "Stores faculty profile and teaching assignment information."),
    ("feedback", "faculty_id, subject, ratings, comments, cycle_id", "Stores anonymous feedback content; it contains no student_id."),
    ("feedback_submissions", "student_id, faculty_id, subject, cycle_id", "Stores participation tracking separately from anonymous response content."),
    ("feedback_cycles", "name, start_date, end_date, status", "Defines feedback windows. The authoritative current cycle is the one marked active."),
    ("hod / admin", "role-specific identifiers and credentials", "Supports HOD and future Admin authentication."),
], [1700, 3250, 4410])

add_callout(doc, "Privacy boundary:", "The feedback table deliberately excludes student_id. Identity-to-participation mapping is stored only in feedback_submissions, so a future checklist can show participation status without exposing feedback answers, comments, response IDs, or submission timestamps.")

doc.add_page_break()
add_heading(doc, "Key Application Workflows", 1)
add_heading(doc, "Student Feedback Submission", 2)
add_numbered(doc, [
    "A student authenticates and receives a server-side student session.",
    "The student selects a faculty member and completes the feedback form.",
    "The backend validates required survey fields and checks for exactly one active feedback cycle.",
    "The backend checks whether the student has already submitted feedback for the same faculty member in that cycle.",
    "Within one database transaction, the backend inserts the anonymous feedback record and the separate participation-tracking record with the same cycle_id.",
    "The transaction commits only if both inserts succeed; otherwise it rolls back."),
)

add_heading(doc, "Faculty Dashboard", 2)
add_bullets(doc, [
    "Loads current faculty profile data using the faculty session.",
    "Uses the feedback cycle whose status is active as the only current-cycle definition.",
    "Retrieves anonymous feedback for the logged-in faculty and active cycle only.",
    "Calculates teaching, communication, behaviour, overall rating, and participation totals.",
    "Returns an explicit no_active_cycle state when no feedback cycle is active instead of using calendar-month inference."),

add_heading(doc, "HOD Analytics", 2)
add_para(doc, "The HOD dashboard aggregates feedback for the authenticated HOD department and presents faculty-level performance metrics. Current HOD reporting remains oriented around department and faculty data; applying explicit cycle filters to HOD reporting is a recommended follow-up for consistent operational reporting.")

add_heading(doc, "Feedback-Cycle Foundation", 1)
add_para(doc, "The project now uses a consistent architectural rule: Current cycle = exactly one feedback_cycles record with status = active. Submission and reporting behavior no longer substitutes calendar-month grouping for a missing active cycle.")
add_table(doc, ["Control", "Implemented Behavior", "Benefit"], [
    ("Historical classification", "Two previously unassigned feedback rows were assigned to the single Legacy Feedback cycle; existing legacy rows were left unchanged.", "All historical feedback now has a valid cycle reference."),
    ("Audit protection", "An audit table records original feedback IDs and original cycle IDs before the one-time legacy backfill.", "Supports traceability and recovery."),
    ("Cycle alignment", "feedback and feedback_submissions use the same cycle_id for every new submission.", "Prevents date-inference ambiguity."),
    ("Duplicate prevention", "Unique student_id + faculty_id + cycle_id constraint.", "Enforces one participation submission per faculty per cycle."),
    ("Referential integrity", "Foreign keys link both feedback tables to feedback_cycles.", "Prevents invalid cycle references."),
    ("No active cycle", "Student submission returns a clear rejection; Faculty API returns no_active_cycle.", "Avoids accidental collection/reporting outside an authorized cycle."),
], [2000, 4500, 2860])

doc.add_page_break()
add_heading(doc, "Security and Privacy Controls", 1)
add_heading(doc, "Implemented Controls", 2)
add_bullets(doc, [
    "Express session management with a MySQL-backed session store and HTTP-only cookies.",
    "Helmet security headers and no-cache response headers for application routes.",
    "Parameterized SQL queries in the inspected route and migration work.",
    "Anonymous feedback records are separate from participation tracking records.",
    "Transactional submission writes prevent a tracking record or anonymous feedback record from being committed alone.",
    "Database uniqueness and foreign keys protect cycle-aware participation integrity."),

add_heading(doc, "Privacy Rules", 2)
add_table(doc, ["Data / Feature", "Allowed", "Not Allowed"], [
    ("Faculty analytics", "Aggregated ratings and anonymous comments for the faculty/cycle.", "Student identity in feedback records or response-level attribution."),
    ("Participation tracking", "Student/faculty/cycle status in the separate tracking table.", "Linking a named student to feedback answers, comments, or feedback IDs."),
    ("Future checklist", "Minimum identity plus Submitted / Not Submitted status.", "Submission date/time, comments, ratings, answers, or any response connection."),
    ("Cycle records", "Historical cycles retained and deactivated.", "Deleting cycles that contain historical feedback."),
], [1900, 3720, 3740])

add_heading(doc, "Known Security Gaps and Constraints", 2)
add_table(doc, ["Area", "Current Finding", "Recommended Resolution"], [
    ("Admin authentication", "The current Admin route is placeholder behavior and no Admin session authorization foundation has been completed.", "Create the first legitimate Admin account locally, then implement database-authenticated Admin login, session regeneration, and require-admin middleware."),
    ("Password scheme", "Student, Faculty, and HOD routes currently appear to use direct password comparison; Admin setup uses salted scrypt hashes for the new setup path.", "Plan a staged, tested credential-hardening migration for all roles; do not mix formats without compatible verification."),
    ("Static Admin pages", "Current static Admin dashboard access needs server-side protection.", "Serve Admin dashboard routes behind require-admin middleware or a protected wrapper."),
    ("HOD cycle scope", "HOD reports are not yet explicitly current-cycle scoped.", "Apply the active-cycle policy when operational reporting requirements demand it."),
], [1750, 4000, 3610])

add_heading(doc, "Testing and Validation Performed", 1)
add_bullets(doc, [
    "Cycle-aware migration executed successfully and re-executed with zero additional legacy rows changed, confirming idempotent behavior.",
    "Database schema verified: feedback.cycle_id and feedback_submissions.cycle_id are non-null and cycle-referenced.",
    "All 18 historical feedback rows verified with valid cycle references.",
    "feedback_submissions confirmed empty after historic migration; no historical identity links were reconstructed.",
    "No active cycle confirmed after migration; the migration intentionally does not create or activate one.",
    "Student submission behavior verified to return NO_ACTIVE_CYCLE before writing data.",
    "Faculty feedback API verified to return a stable no_active_cycle response without crashing.",
    "Modified backend and migration scripts passed JavaScript syntax validation."),

doc.add_page_break()
add_heading(doc, "Current Project Status", 1)
add_table(doc, ["Workstream", "Status", "Notes"], [
    ("Anonymous feedback submission", "Implemented", "Cycle-aware transactional submission is in place; submissions require one active cycle."),
    ("Faculty dashboard", "Implemented", "Professional analytics dashboard; current-cycle data requires an active cycle."),
    ("HOD dashboard", "Implemented", "Department and faculty analytics available; cycle scoping is a future improvement."),
    ("Feedback cycle data foundation", "Implemented", "Cycle integrity, tracking alignment, historic classification, and no-fallback behavior completed."),
    ("Admin account bootstrap", "Implemented setup path", "Local create-admin.js creates one legitimate hashed Admin account; it must be run manually."),
    ("Admin authentication and authorization", "Pending", "Requires a legitimate Admin account before secure login/session middleware is implemented and tested."),
    ("Admin cycle management", "Planned", "Must be protected by Admin-only backend authorization."),
    ("Faculty submission checklist", "Planned", "Must use only separate participation tracking and omit timestamps/feedback content."),
], [2400, 1750, 5210])

add_heading(doc, "Recommended Next Steps", 1)
add_numbered(doc, [
    "Run create-admin.js locally to create the first legitimate Admin account without sharing credentials through chat.",
    "Implement secure Admin login, session creation, server-side require-admin middleware, protected Admin dashboard delivery, and logout.",
    "Build Admin-only feedback-cycle management: list, create, activate, and deactivate cycles while preventing multiple active cycles and preserving historical cycles.",
    "Test a full active-cycle lifecycle: create cycle, activate, submit feedback, reject duplicates, deactivate, and verify dashboard state.",
    "Add the Faculty participation checklist using students plus feedback_submissions only; expose identity and Submitted/Not Submitted status without feedback data or timestamps.",
    "Extend HOD reports with explicit active-cycle and historical-cycle filters as required.",
    "Plan a controlled credential-security upgrade for existing Student, Faculty, and HOD password storage."),

add_heading(doc, "Conclusion", 1)
add_para(doc, "The project has a strong functional foundation for anonymous feedback collection and faculty analytics. Its most important recent improvement is the separation and alignment of anonymous response data with cycle-aware participation tracking. The next critical milestone is a secure Admin authentication layer, followed by protected cycle management. Completing those steps will allow the system to operate feedback windows predictably while retaining privacy and preserving historical records.")

doc.save(OUTPUT)
print(OUTPUT)
